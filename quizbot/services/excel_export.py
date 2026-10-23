"""Administrative data snapshots and formatted openpyxl exports."""

import asyncio
import io
import os
import re
import tempfile
from collections import Counter
from datetime import UTC
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select, text

from quizbot.db.models import (
    AnswerKey,
    Attempt,
    Exam,
    GradingResult,
    Question,
    QuestionPart,
    StudentResponse,
    User,
)
from quizbot.domain import DomainError
from quizbot.services.base import Service
from quizbot.services.simple_grading import element_from_row, summarize


def excel_value(value):
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
    return value


def local_timestamp(value, timezone: str) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(ZoneInfo(timezone)).isoformat(timespec="seconds")


class ExcelExportService(Service):
    async def simple_report(self, actor: int, exam_id: int) -> dict:
        self.admin(actor)
        async with self.sessions.begin() as session:
            if session.bind.dialect.name == "postgresql":
                await session.execute(
                    text("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
                )
            exam = await session.get(Exam, exam_id)
            if exam is None:
                raise DomainError("Test topilmadi.")
            attempts = (
                await session.execute(
                    select(Attempt, User)
                    .join(User)
                    .where(Attempt.exam_id == exam_id)
                    .order_by(Attempt.id)
                )
            ).all()
            definitions = (
                await session.execute(
                    select(Question, QuestionPart, AnswerKey)
                    .join(QuestionPart, QuestionPart.question_id == Question.id)
                    .join(AnswerKey, AnswerKey.part_id == QuestionPart.id)
                    .where(Question.exam_id == exam_id)
                    .order_by(Question.number, QuestionPart.label)
                )
            ).all()
            responses = {
                (r.attempt_id, r.part_id): r
                for r in await session.scalars(
                    select(StudentResponse).join(Attempt).where(Attempt.exam_id == exam_id)
                )
            }
            grades = {
                (g.attempt_id, g.part_id): g
                for g in await session.scalars(
                    select(GradingResult).join(Attempt).where(Attempt.exam_id == exam_id)
                )
            }
            participants, answers, outcomes = [], [], []
            question_states = {q.id: Counter() for q, _p, _k in definitions}
            for attempt, user in attempts:
                elements = [
                    element_from_row(
                        q, p, k, responses.get((attempt.id, p.id)), grades.get((attempt.id, p.id))
                    )
                    for q, p, k in definitions
                ]
                metrics = summarize(elements)
                submitted = attempt.status == "submitted"
                if submitted:
                    outcomes.append(metrics)
                    for qid, state in metrics["parent_states"].items():
                        question_states[qid][state] += 1
                participant = {
                    "student_id": user.id,
                    "attempt_id": attempt.id,
                    "telegram_id": user.telegram_id,
                    "first_name": user.first_name,
                    "last_name": user.last_name,
                    "full_name": user.full_name,
                    "exam_id": exam_id,
                    "exam_title": exam.title,
                    "status": attempt.status,
                    "submission_date": local_timestamp(
                        attempt.submitted_at, self.settings.report_timezone
                    ),
                    "manual_review_status": "Topshirilmagan"
                    if not submitted
                    else "Tekshirilmoqda"
                    if metrics["pending_elements"]
                    else "Tekshirilgan",
                    **{k: v for k, v in metrics.items() if k != "parent_states"},
                }
                if not submitted:
                    for key in (
                        "correct_questions",
                        "incorrect_questions",
                        "partial_questions",
                        "unanswered_questions",
                        "correct_elements",
                        "incorrect_elements",
                        "raw_score",
                        "percentage",
                    ):
                        participant[key] = None
                participants.append(participant)
                if submitted:
                    answers.extend(
                        {
                            "student_id": user.id,
                            "full_name": user.full_name,
                            "attempt_id": attempt.id,
                            "question_number": a.number,
                            "part": a.label,
                            "student_answer": a.answer,
                            "correct_answer": a.accepted[0] if a.accepted else "Qo‘lda baholanadi",
                            "is_correct": None if a.state == "pending" else a.state == "correct",
                            "state": a.state,
                            "awarded_points": a.points,
                            "maximum_points": a.maximum,
                        }
                        for a in elements
                    )
            question_numbers = {q.id: q.number for q, _p, _k in definitions}
            items = []
            for qid, counts in question_states.items():
                resolved = sum(counts.values()) - counts["pending"]
                items.append(
                    {
                        "question": question_numbers[qid],
                        **{
                            state: counts[state]
                            for state in (
                                "correct",
                                "incorrect",
                                "partial",
                                "unanswered",
                                "pending",
                            )
                        },
                        "responses": sum(counts.values()),
                        "observed_percent_correct": 100 * counts["correct"] / resolved
                        if resolved
                        else None,
                    }
                )
            items.sort(key=lambda r: r["question"])
            ranked = sorted(
                [r for r in items if r["observed_percent_correct"] is not None],
                key=lambda r: (r["observed_percent_correct"], r["question"]),
            )
            complete = [m for m in outcomes if not m["pending_elements"]]

            def average(key):
                return sum(m[key] for m in complete) / len(complete) if complete else None

            summary = {
                "exam_id": exam_id,
                "title": exam.title,
                "subject": exam.subject,
                "participants": len(attempts),
                "submitted": len(outcomes),
                "pending_manual_review": sum(m["pending_elements"] > 0 for m in outcomes),
                "pending_elements": sum(m["pending_elements"] for m in outcomes),
                "average_correct_questions": average("correct_questions"),
                "average_raw": average("raw_score"),
                "average_percentage": average("percentage"),
                "total_questions": len(question_states),
                "total_elements": len(definitions),
                "timezone": self.settings.report_timezone,
                "release_policy": exam.simple_release_policy or self.settings.simple_release_policy,
                "results_released": exam.simple_results_released,
                "answer_review_enabled": exam.answer_review_enabled,
                "hardest": ranked[:5],
                "easiest": list(reversed(ranked[-5:])),
                # Kept only for internal compatibility; never rendered in simple mode.
                "average_t": None,
                "grades": {},
                "calibration_status": "disabled",
                "strategy": "simple",
            }
            return {
                "summary": summary,
                "participants": participants,
                "answers": answers,
                "items": items,
            }

    async def export_simple_excel(self, actor: int, exam_id: int) -> bytes:
        self.admin(actor)
        report = await self.simple_report(actor, exam_id)
        return await asyncio.to_thread(
            write_workbook, report, self.settings.results_dir / f"exam_{exam_id}_results.xlsx"
        )


def format_sheet(sheet, columns):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="14532D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 34
    for index, _column in enumerate(columns, 1):
        values = [
            str(sheet.cell(row, index).value or "") for row in range(1, min(sheet.max_row, 200) + 1)
        ]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            55, max(14, max(map(len, values), default=14) + 2)
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0FDF4")


def build_workbook(report: dict) -> bytes:
    book = Workbook()
    book.remove(book.active)
    columns = [
        ("Student ID", "student_id"),
        ("Telegram ID", "telegram_id"),
        ("First Name", "first_name"),
        ("Last Name", "last_name"),
        ("Exam ID", "exam_id"),
        ("Exam Title", "exam_title"),
        ("Correct Questions", "correct_questions"),
        ("Incorrect Questions", "incorrect_questions"),
        ("Partially Correct Questions", "partial_questions"),
        ("Unanswered Questions", "unanswered_questions"),
        ("Correct Scoring Elements", "correct_elements"),
        ("Incorrect Scoring Elements", "incorrect_elements"),
        ("Raw Score", "raw_score"),
        ("Maximum Raw Score", "maximum_raw_score"),
        ("Percentage", "percentage"),
        ("Submission Date", "submission_date"),
        ("Manual Review Status", "manual_review_status"),
    ]
    answer_columns = [
        ("Student ID", "student_id"),
        ("Student Full Name", "full_name"),
        ("Question Number", "question_number"),
        ("Question Part", "part"),
        ("Student Answer", "student_answer"),
        ("Correct Answer", "correct_answer"),
        ("Is Correct", "is_correct"),
        ("Awarded Points", "awarded_points"),
        ("Maximum Points", "maximum_points"),
    ]
    for name, fields, data in [
        ("Students", columns, report["participants"]),
        ("Individual Answers", answer_columns, report["answers"]),
    ]:
        sheet = book.create_sheet(name)
        sheet.append([label for label, _key in fields])
        for row in data:
            sheet.append([excel_value(row.get(key)) for _label, key in fields])
        format_sheet(sheet, fields)
    summary = book.create_sheet("Exam Summary")
    summary.append(
        [
            "Section",
            "Question Number",
            "Correct",
            "Incorrect",
            "Partial",
            "Unanswered",
            "Pending",
            "Responses",
            "Observed Correct %",
            "Metric",
            "Value",
        ]
    )
    for metric in (
        "exam_id",
        "title",
        "timezone",
        "participants",
        "submitted",
        "pending_manual_review",
        "average_correct_questions",
        "average_raw",
        "average_percentage",
    ):
        summary.append(
            [
                "Overall",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                metric,
                excel_value(report["summary"][metric]),
            ]
        )
    for section, data in [
        ("Question statistics", report["items"]),
        ("Most difficult (observed)", report["summary"]["hardest"]),
        ("Easiest (observed)", report["summary"]["easiest"]),
    ]:
        for row in data:
            summary.append(
                [
                    section,
                    row["question"],
                    row["correct"],
                    row["incorrect"],
                    row["partial"],
                    row["unanswered"],
                    row["pending"],
                    row["responses"],
                    row["observed_percent_correct"],
                ]
            )
    format_sheet(summary, summary[1])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def write_workbook(report: dict, path: Path) -> bytes:
    payload = build_workbook(report)
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=path.parent, prefix=f".{path.stem}-", suffix=".xlsx", delete=False
        ) as stream:
            temporary = Path(stream.name)
            stream.write(payload)
        os.replace(temporary, path)
    finally:
        if temporary and temporary.exists():
            temporary.unlink()
    return payload
