import io

import pytest
from openpyxl import load_workbook

from quizbot.bot.admin import perform
from quizbot.bot.student import show_result, show_review
from quizbot.domain import Forbidden
from quizbot.importing import ExamSpec

pytestmark = pytest.mark.postgres


class Message:
    def __init__(self):
        self.messages = []
        self.photos = []
        self.documents = []

    async def answer(self, text, **kwargs):
        self.messages.append({"text": text, **kwargs})

    async def answer_photo(self, photo):
        self.photos.append(photo)

    async def answer_document(self, document):
        self.documents.append(document)


async def setup(platform):
    exam = await platform.import_exam(
        999,
        ExamSpec(
            title="Ona tili <test>",
            subject="Ona tili",
            materials=[{"key": "read", "title": "Asl matn", "text": "Test o‘qish matni"}],
            questions=[
                {
                    "number": 40,
                    "kind": "multipart",
                    "statement": "Asl savol <tag>",
                    "image_file_id": "original-diagram",
                    "material_keys": ["read"],
                    "rubric": {"explanation": "Asl tushuntirish"},
                    "parts": [
                        {"label": "A", "accepted": ["x"]},
                        {"label": "B", "accepted": ["=Correct()"]},
                    ],
                }
            ],
        ),
    )
    await platform.set_published(999, exam.id, True)
    await platform.register(123, first_name="=Ali", last_name="<Karimov>")
    attempt = await platform.start_attempt(123, exam.id)
    for index, answer in [(0, "x"), (1, "=Student()")]:
        view = await platform.question_view(123, attempt.id, index)
        await platform.save_answer(123, attempt.id, view["part_id"], answer)
    await platform.submit_attempt(123, attempt.id)
    return exam, attempt


async def test_simple_student_result_review_and_admin_stats_use_formatted_messages(simple_platform):
    exam, attempt = await setup(simple_platform)
    message = Message()
    await show_result(message, simple_platform, 123, attempt.id)
    result = message.messages[-1]
    assert result["parse_mode"] == "HTML"
    assert "MILLIY SERTIFIKAT MOCK" in result["text"]
    assert "Qisman to‘g‘ri: <b>1 ta" in result["text"]
    buttons = [b.text for row in result["reply_markup"].inline_keyboard for b in row]
    assert buttons == [
        "❌ Xato javoblarni ko‘rish",
        "✅ To‘g‘ri javoblarni ko‘rish",
        "🏠 Bosh menyu",
    ]
    await show_review(message, simple_platform, 123, attempt.id, "wrong", 0)
    assert message.photos == ["original-diagram"]
    text = message.messages[-1]["text"]
    assert "B qism" in text and "Asl savol &lt;tag&gt;" in text
    assert "=Student()" in text and "=Correct()" in text and "Asl tushuntirish" in text
    assert any(
        b.text == "📖 O‘qish matni"
        for row in message.messages[-1]["reply_markup"].inline_keyboard
        for b in row
    )
    assert (await simple_platform.review_materials(123, attempt.id, "wrong"))[0][
        "text"
    ] == "Test o‘qish matni"
    await perform(message, simple_platform, 999, "stats", exam.id)
    stats = message.messages[-1]
    assert stats["parse_mode"] == "HTML" and "TEST STATISTIKASI" in stats["text"]
    assert "Ona tili &lt;test&gt;" in stats["text"]
    assert all(
        word not in stats["text"] for word in ("calibration_status", "average_t", "{", "null")
    )
    with pytest.raises(Forbidden):
        await show_review(message, simple_platform, 456, attempt.id, "wrong", 0)


async def test_new_excel_renderer_escapes_formulas_in_names_answers_and_keys(simple_platform):
    exam, _attempt = await setup(simple_platform)
    payload = await simple_platform.export_report(999, exam.id)
    book = load_workbook(io.BytesIO(payload))
    assert book["Students"].cell(2, 3).value == "'=Ali"
    answers = book["Individual Answers"]
    assert answers.cell(3, 5).value == "'=Student()"
    assert answers.cell(3, 6).value == "'=Correct()"
    assert answers.cell(3, 5).data_type == "s"
    with pytest.raises(Forbidden):
        await simple_platform.simple_report(123, exam.id)
    with pytest.raises(Forbidden):
        await simple_platform.review_materials(456, _attempt.id, "wrong")
