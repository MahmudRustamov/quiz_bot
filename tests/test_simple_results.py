import io
from dataclasses import replace
from datetime import UTC, datetime
from html.parser import HTMLParser

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from quizbot.config import Settings
from quizbot.db.models import (
    Attempt,
    CalibrationRun,
    Exam,
    StandardizedScore,
    StudentAbilityEstimate,
    StudentResponse,
)
from quizbot.domain import DomainError, Forbidden
from quizbot.importing import ExamSpec
from quizbot.services.excel_export import local_timestamp
from quizbot.services.result_presentation import ResultPresentationService
from quizbot.services.simple_grading import ElementResult, summarize


def element(question_id, state, label="", points=None, maximum=1.0):
    return ElementResult(
        question_id,
        question_id,
        question_id * 10 + int(label == "B"),
        label,
        state,
        None if state == "unanswered" else "x",
        ["x"],
        points if points is not None else float(state == "correct"),
        maximum,
    )


def test_parent_counts_do_not_infer_correctness_from_raw_score():
    stats = summarize(
        [
            element(1, "correct", points=14, maximum=14),
            element(2, "correct", "A"),
            element(2, "incorrect", "B"),
            element(3, "incorrect"),
            element(4, "unanswered"),
        ]
    )
    assert stats["raw_score"] == 15
    assert stats["correct_questions"] == 1
    assert stats["incorrect_questions"] == 1
    assert stats["partial_questions"] == 1
    assert stats["unanswered_questions"] == 1
    assert stats["correct_elements"] == 2
    assert stats["incorrect_elements"] == 2
    assert stats["unanswered_elements"] == 1
    assert stats["percentage"] == 25
    assert stats["element_percentage"] == 40


@pytest.mark.parametrize(
    "states,expected",
    [
        (("correct", "correct"), "correct"),
        (("correct", "incorrect"), "partial"),
        (("correct", "unanswered"), "partial"),
        (("incorrect", "unanswered"), "incorrect"),
        (("unanswered", "unanswered"), "unanswered"),
        (("correct", "pending"), "pending"),
    ],
)
def test_multipart_classification(states, expected):
    stats = summarize([element(40, states[0], "A"), element(40, states[1], "B")])
    assert stats[f"{expected}_questions"] == 1
    assert (
        sum(
            stats[f"{s}_questions"]
            for s in ("correct", "incorrect", "partial", "unanswered", "pending")
        )
        == 1
    )
    if expected == "pending":
        assert stats["raw_score"] is None


def result_data():
    return {
        "full_name": "Ali <Karimov>",
        "subject": "Ona tili & adabiyot",
        "title": "19.10 — MS Mock",
        "release_state": "released",
        "metrics": summarize(
            [element(1, "correct"), element(2, "incorrect"), element(3, "unanswered")]
        ),
    }


def test_simple_presentation_escapes_html_and_has_no_psychometric_fields():
    result = result_data()
    result.update(theta=2, t=70, grade="A+", calibrated=True)
    message = ResultPresentationService.student(result)
    assert "Ali &lt;Karimov&gt;" in message
    assert "Ona tili &amp; adabiyot" in message
    assert "33.3%" in message
    assert "<Karimov>" not in message
    assert all(word not in message for word in ("θ", "Z-score", "T-score", "A+", "Kalibrlash"))


@pytest.mark.parametrize(
    "state,expected",
    [("reviewing", "Tekshirilmoqda"), ("held", "Administrator"), ("in_progress", "tugallanmagan")],
)
def test_unreleased_presentation_never_leaks_scores(state, expected):
    result = result_data()
    result["release_state"] = state
    message = ResultPresentationService.student(result)
    assert expected in message
    assert "33.3%" not in message and "Xom ball" not in message


def test_review_splits_long_html_and_preserves_question_and_answers():
    a = replace(
        element(40, "incorrect", "<A>"),
        statement="<>& " * 800,
        answer="<bad>",
        accepted=["<right>"],
        explanation="& " * 2000,
    )
    messages = ResultPresentationService.review(
        {"answer": a, "index": 0, "count": 1, "options": []}
    )
    assert len(messages) > 1 and all(len(m) <= 3900 for m in messages)
    joined = "\n".join(messages)
    assert "&lt;bad&gt;" in joined and "&lt;right&gt;" in joined and "&lt;A&gt;" in joined

    class Parser(HTMLParser):
        def handle_starttag(self, tag, attrs):
            assert tag in {"b", "i"}

    for message in messages:
        Parser().feed(message)


def test_timezone_is_explicit_and_configurable():
    assert (
        local_timestamp(datetime(2026, 10, 23, 9, tzinfo=UTC), "Asia/Tashkent")
        == "2026-10-23T14:00:00+05:00"
    )
    with pytest.raises(ValueError):
        Settings(_env_file=None, report_timezone="Invalid/Timezone")


def simple_spec(*, manual=False, release=None):
    return ExamSpec(
        title="Actual exam",
        subject="Ona tili va adabiyot",
        simple_release_policy=release,
        questions=[
            {
                "number": 1,
                "kind": "mcq",
                "statement": "Actual question",
                "options": {"A": "First", "B": "Second"},
                "parts": [{"accepted": ["A"], "weight": 14}],
            },
            {
                "number": 40,
                "kind": "multipart",
                "statement": "Actual multipart",
                "parts": [
                    {"label": "A", "accepted": ["x"], "manual_review": manual},
                    {"label": "B", "accepted": ["y"]},
                ],
            },
            {
                "number": 41,
                "kind": "written",
                "statement": "Actual blank",
                "parts": [{"accepted": ["z"]}],
            },
        ],
    )


async def prepare(platform, *, manual=False, release=None):
    exam = await platform.import_exam(999, simple_spec(manual=manual, release=release))
    await platform.set_published(999, exam.id, True)
    await platform.register(123, first_name="Ali", last_name="Karimov")
    attempt = await platform.start_attempt(123, exam.id)
    for index, answer in [(0, "A"), (1, "x"), (2, "wrong")]:
        view = await platform.question_view(123, attempt.id, index)
        await platform.save_answer(123, attempt.id, view["part_id"], answer)
    return exam, attempt


@pytest.mark.postgres
async def test_simple_immediate_results_with_one_student_and_no_calibration(simple_platform):
    platform = simple_platform
    exam, attempt = await prepare(platform)
    await platform.submit_attempt(123, attempt.id)
    result = await platform.result(123, attempt.id)
    assert result["released"] and result["review_enabled"]
    assert result["metrics"]["correct_questions"] == 1
    assert result["metrics"]["partial_questions"] == 1
    assert result["metrics"]["unanswered_questions"] == 1
    assert result["metrics"]["percentage"] == pytest.approx(100 / 3)
    assert result["raw"] == 15 and result["maximum"] == 17
    assert not any(key in result for key in ("theta", "z", "t", "grade"))
    async with platform.sessions() as session:
        assert not (await session.get(Exam, exam.id)).results_released
        assert not list(await session.scalars(select(CalibrationRun)))
    with pytest.raises(DomainError, match="o‘chirilgan"):
        await platform.run_calibration(999, exam.id)


@pytest.mark.postgres
async def test_manual_release_is_independent_of_calibration_and_closed_cohort(simple_platform):
    platform = simple_platform
    exam, attempt = await prepare(platform, release="manual")
    await platform.submit_attempt(123, attempt.id)
    assert (await platform.result(123, attempt.id))["release_state"] == "held"
    with pytest.raises(DomainError):
        await platform.review_page(123, attempt.id, "wrong")
    await platform.release_results(999, exam.id)
    assert (await platform.result(123, attempt.id))["released"]
    async with platform.sessions() as session:
        stored = await session.get(Exam, exam.id)
        assert not stored.closed and not stored.results_released and stored.simple_results_released


@pytest.mark.postgres
async def test_pending_manual_grade_remains_hidden_even_after_admin_release(simple_platform):
    platform = simple_platform
    exam, attempt = await prepare(platform, manual=True, release="manual")
    await platform.submit_attempt(123, attempt.id)
    await platform.release_results(999, exam.id)
    result = await platform.result(123, attempt.id)
    assert result["release_state"] == "reviewing" and result["metrics"] is None
    pending = await platform.pending_reviews(999, exam.id)
    await platform.review_grade(999, pending[0]["id"], 1, "Checked")
    assert (await platform.result(123, attempt.id))["released"]
    with pytest.raises(DomainError):
        await platform.review_grade(999, pending[0]["id"], 0, "Cannot change a released result")


@pytest.mark.postgres
async def test_review_authorization_filters_and_reading_materials(simple_platform):
    platform = simple_platform
    exam, attempt = await prepare(platform)
    with pytest.raises(DomainError):
        await platform.review_page(123, attempt.id, "correct")
    await platform.submit_attempt(123, attempt.id)
    wrong = await platform.review_page(123, attempt.id, "wrong")
    assert wrong["count"] == 1 and wrong["answer"].number == 40 and wrong["answer"].label == "B"
    assert wrong["answer"].accepted == ["y"] and wrong["answer"].answer == "wrong"
    assert (await platform.review_page(123, attempt.id, "correct"))["count"] == 2
    with pytest.raises(Forbidden):
        await platform.review_page(456, attempt.id, "wrong")
    with pytest.raises(Forbidden):
        await platform.export_report(123, exam.id)
    await platform.set_result_policy(999, exam.id, "immediate", review_enabled=False)
    assert not (await platform.result(123, attempt.id))["review_enabled"]
    with pytest.raises(DomainError):
        await platform.review_page(123, attempt.id, "wrong")


@pytest.mark.postgres
async def test_excel_contains_actual_counts_answers_and_timezone(simple_platform):
    platform = simple_platform
    exam, attempt = await prepare(platform)
    await platform.submit_attempt(123, attempt.id)
    async with platform.sessions.begin() as session:
        await session.execute(
            update(Attempt)
            .where(Attempt.id == attempt.id)
            .values(submitted_at=datetime(2026, 10, 23, 9, tzinfo=UTC))
        )
    payload = await platform.export_report(999, exam.id)
    path = platform.settings.results_dir / f"exam_{exam.id}_results.xlsx"
    assert path.read_bytes() == payload
    book = load_workbook(io.BytesIO(payload))
    assert book.sheetnames == ["Students", "Individual Answers", "Exam Summary"]
    students = book["Students"]
    row = dict(zip([c.value for c in students[1]], [c.value for c in students[2]], strict=True))
    assert row["First Name"] == "Ali" and row["Last Name"] == "Karimov"
    assert row["Correct Questions"] == 1 and row["Partially Correct Questions"] == 1
    assert row["Correct Scoring Elements"] == 2
    assert row["Raw Score"] == 15 and row["Percentage"] == pytest.approx(100 / 3)
    assert row["Submission Date"] == "2026-10-23T14:00:00+05:00"
    assert students.freeze_panes == "A2" and students.auto_filter.ref
    assert book["Individual Answers"].max_row == 5
    summary = (await platform.report(999, exam.id))["summary"]
    assert summary["average_correct_questions"] == 1 and summary["average_raw"] == 15
    assert "observed_percent_correct" in summary["hardest"][0]


@pytest.mark.postgres
async def test_simple_mode_preserves_raw_responses_and_future_rasch_data(simple_platform):
    platform = simple_platform
    exam, attempt = await prepare(platform)
    await platform.submit_attempt(123, attempt.id)
    async with platform.sessions.begin() as session:
        run = CalibrationRun(
            exam_id=exam.id,
            status="ready",
            model="Rasch-1PL",
            algorithm_version="historical",
            snapshot={},
        )
        session.add(run)
        await session.flush()
        session.add(
            StudentAbilityEstimate(
                run_id=run.id, attempt_id=attempt.id, theta=1.0, posterior_sd=0.5
            )
        )
        session.add(
            StandardizedScore(
                run_id=run.id, attempt_id=attempt.id, z_score=1.0, t_score=60.0, grade="B+"
            )
        )
        stored_exam = await session.get(Exam, exam.id)
        stored_exam.released_run_id = run.id
        stored_exam.results_released = True
    result = await platform.result(123, attempt.id)
    assert not result["calibrated"] and "theta" not in result and result["raw"] == 15
    async with platform.sessions() as session:
        assert (await session.get(Attempt, attempt.id)).raw_score == 15
        assert len(list(await session.scalars(select(StudentResponse)))) == 3
        assert (await session.scalar(select(StandardizedScore))).t_score == 60


@pytest.mark.postgres
async def test_registration_requires_explicit_names_and_retains_timestamp(simple_platform):
    platform = simple_platform
    with pytest.raises(DomainError):
        await platform.register(123, "OnlyFirst")
    first = await platform.register(123, first_name="Ali", last_name="Karimov")
    second = await platform.register(123, first_name="Vali", last_name="Rahimov")
    assert first.created_at == second.created_at
    assert second.first_name == "Vali" and second.last_name == "Rahimov"
